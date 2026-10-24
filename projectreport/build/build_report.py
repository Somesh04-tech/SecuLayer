"""Build projectreport/SecuLayer_Project_Report.docx from report_content.py.

The document starts from the college format (projectreport/MDM_Project_report (1).docx, read
only): its page size, theme and logo are reused, and its styles are adjusted to the format's
note (body 12 pt Times New Roman at 1.5 line spacing, headings 14 pt bold, sub-headings 14 pt).

Structure: section 1 (cover, certificate: no page numbers), section 2 (front matter: lower-case
roman page numbers), section 3 (chapters, references, appendices: arabic from 1, with a running
header). The contents, list of figures and list of tables are real TOC fields whose cached
entries carry page numbers measured from a LibreOffice PDF render (two passes), so they read
correctly before anyone presses "Update field" in Word.

Usage (from the repository root):
  uv run --no-project --with python-docx --with pypdf python projectreport/build/build_report.py
"""
from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import report_content as C  # noqa: E402

PR = HERE.parent
FORMAT = PR / "MDM_Project_report (1).docx"
FIGS = PR / "figures"
OUT = PR / "SecuLayer_Project_Report.docx"
WORK = HERE / "out"
TEXT_WIDTH = 6.5  # inches: Letter (8.5 in) less two 1-inch margins, as in the format
MAX_FIG_H = 7.4
FONT = "Times New Roman"
MONO = "Courier New"
BLACK = RGBColor(0, 0, 0)


# ------------------------------------------------------------------------------------------------
# Low-level XML helpers
def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para_shading(paragraph, hex_fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def para_border_bottom(paragraph, sz=6, color="808080"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def fld_char(run, kind):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._r.append(el)


def instr(run, text):
    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = text
    run._r.append(el)


def add_simple_field(paragraph, code, cached, bold=None, size=None):
    r = paragraph.add_run()
    fld_char(r, "begin")
    r = paragraph.add_run()
    instr(r, f" {code} ")
    r = paragraph.add_run()
    fld_char(r, "separate")
    r = paragraph.add_run(cached)
    if bold is not None:
        r.bold = bold
    if size is not None:
        r.font.size = size
    r = paragraph.add_run()
    fld_char(r, "end")


def set_rfonts(style_or_run_element, name):
    rPr = style_or_run_element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rFonts.get(qn(attr)) is not None:
            del rFonts.attrib[qn(attr)]


def outline_level(style, level):
    pPr = style.element.get_or_add_pPr()
    for old in pPr.findall(qn("w:outlineLvl")):
        pPr.remove(old)
    if level is not None:
        el = OxmlElement("w:outlineLvl")
        el.set(qn("w:val"), str(level))
        pPr.append(el)


def roman(n):
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


# ------------------------------------------------------------------------------------------------
# Document setup
def fresh_document():
    doc = Document(str(FORMAT))
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    sect = doc.sections[0]
    sect.page_width, sect.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sect, side, Inches(1.0))
    sect.header_distance = Inches(0.5)
    sect.footer_distance = Inches(0.5)
    sect.start_type = WD_SECTION_START.NEW_PAGE
    setup_styles(doc)
    setup_numbering(doc)
    return doc


def style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, kind)
        if base and kind == WD_STYLE_TYPE.PARAGRAPH:
            st.base_style = doc.styles[base]
    return st


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    set_rfonts(normal.element, FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    def heading(name, size, bold, align, before, after, level, page_break=False):
        st = style(doc, name)
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = False
        st.font.color.rgb = BLACK
        set_rfonts(st.element, FONT)
        f = st.paragraph_format
        f.alignment = align
        f.space_before = Pt(before)
        f.space_after = Pt(after)
        f.keep_with_next = True
        f.page_break_before = page_break
        f.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        f.left_indent = Inches(0)
        f.first_line_indent = Inches(0)
        outline_level(st, level)
        return st

    # The format's note: headings 14 pt bold; sub-headings 14 pt. Third level: 12 pt bold.
    heading("Heading 1", 14, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 18, 0, page_break=True)
    heading("Heading 2", 14, False, WD_ALIGN_PARAGRAPH.LEFT, 12, 6, 1)
    heading("Heading 3", 12, True, WD_ALIGN_PARAGRAPH.LEFT, 10, 4, 2)
    heading("Front Heading", 14, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 18, 0)
    heading("Contents Heading", 14, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 18, None)

    cap = style(doc, "Caption")
    cap.font.name = FONT
    cap.font.size = Pt(11)
    cap.font.bold = False
    cap.font.italic = False
    cap.font.color.rgb = BLACK
    set_rfonts(cap.element, FONT)
    f = cap.paragraph_format
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.line_spacing_rule = WD_LINE_SPACING.SINGLE
    f.space_before = Pt(4)
    f.space_after = Pt(12)

    for name, size, bold, indent in (("TOC 1", 12, True, 0.0), ("TOC 2", 12, False, 0.35), ("TOC 3", 11, False, 0.7),
                                     ("Table of Figures", 12, False, 0.0)):
        st = style(doc, name)
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        set_rfonts(st.element, FONT)
        f = st.paragraph_format
        f.alignment = WD_ALIGN_PARAGRAPH.LEFT
        f.line_spacing_rule = WD_LINE_SPACING.SINGLE
        f.space_before = Pt(6 if name == "TOC 1" else 2)
        f.space_after = Pt(2)
        f.left_indent = Inches(indent + (0.9 if name == "Table of Figures" else 0))
        f.first_line_indent = Inches(-0.9 if name == "Table of Figures" else 0)
        f.tab_stops.add_tab_stop(Inches(TEXT_WIDTH), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    tt = style(doc, "Table Text")
    tt.font.name = FONT
    tt.font.size = Pt(10.5)
    set_rfonts(tt.element, FONT)
    f = tt.paragraph_format
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f.line_spacing_rule = WD_LINE_SPACING.SINGLE
    f.space_before = Pt(1)
    f.space_after = Pt(1)

    code = style(doc, "Code Block")
    code.font.name = MONO
    code.font.size = Pt(8.5)
    set_rfonts(code.element, MONO)
    f = code.paragraph_format
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f.line_spacing_rule = WD_LINE_SPACING.SINGLE
    f.space_before = Pt(0)
    f.space_after = Pt(0)
    f.left_indent = Inches(0.15)
    f.keep_together = True
    f.keep_with_next = True

    for name in ("List Bullet Custom", "List Number Custom"):
        st = style(doc, name)
        f = st.paragraph_format
        f.space_after = Pt(3)
        f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    ref = style(doc, "Reference Entry")
    f = ref.paragraph_format
    f.left_indent = Inches(0.45)
    f.first_line_indent = Inches(-0.45)
    f.line_spacing_rule = WD_LINE_SPACING.SINGLE
    f.space_after = Pt(6)
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ("Header", "Footer"):
        st = style(doc, name)
        st.font.name = FONT
        st.font.size = Pt(10)
        set_rfonts(st.element, FONT)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        st.paragraph_format.space_after = Pt(0)


NUM_STATE = {}


def setup_numbering(doc):
    numbering = doc.part.numbering_part.element
    ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    nums = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    next_abs = max(ids + [0]) + 1
    NUM_STATE["numbering"] = numbering
    NUM_STATE["next_num"] = max(nums + [0]) + 1

    def abstract(fmt, text, font=None):
        nonlocal next_abs
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(next_abs))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "singleLevel")
        a.append(mlt)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:lvlText", text), ("w:lvlJc", "left")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            lvl.append(e)
        pPr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "360")
        pPr.append(ind)
        lvl.append(pPr)
        if font:
            rPr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for k in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(k), font)
            rPr.append(rf)
            lvl.append(rPr)
        a.append(lvl)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        else:
            numbering.append(a)
        next_abs += 1
        return next_abs - 1

    NUM_STATE["bullet_abs"] = abstract("bullet", "•", FONT)
    NUM_STATE["decimal_abs"] = abstract("decimal", "%1.")
    NUM_STATE["bullet_num"] = new_num(NUM_STATE["bullet_abs"])


def new_num(abstract_id, restart=False):
    numbering = NUM_STATE["numbering"]
    n = OxmlElement("w:num")
    nid = NUM_STATE["next_num"]
    NUM_STATE["next_num"] += 1
    n.set(qn("w:numId"), str(nid))
    a = OxmlElement("w:abstractNumId")
    a.set(qn("w:val"), str(abstract_id))
    n.append(a)
    if restart:
        o = OxmlElement("w:lvlOverride")
        o.set(qn("w:ilvl"), "0")
        s = OxmlElement("w:startOverride")
        s.set(qn("w:val"), "1")
        o.append(s)
        n.append(o)
    numbering.append(n)
    return nid


def set_num(paragraph, num_id):
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), str(num_id))
    numPr.append(il)
    numPr.append(ni)
    pPr.append(numPr)


# ------------------------------------------------------------------------------------------------
# Inline markup: **bold**, *italic*, `code`
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")


def add_rich(paragraph, text, size=None, bold=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = MONO
            set_rfonts(r._r, MONO)
            r.font.size = Pt((size.pt if size else 12) - 1.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        else:
            r = paragraph.add_run(part)
        if size is not None and not part.startswith("`"):
            r.font.size = size
        if bold is not None and not part.startswith("**"):
            r.bold = bold if r.bold is None else r.bold


# ------------------------------------------------------------------------------------------------
# Numbering of chapters, sections, figures, tables and citations
def plan(blocks):
    """Assign numbers. Returns (resolved blocks, figures, tables, refs order)."""
    figs, tabs, order = {}, {}, []
    chap, sec, sub, fn, tn = "", 0, 0, 0, 0
    out = []
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            _, label, title = b
            chap, sec, sub, fn, tn = label, 0, 0, 0, 0
            out.append(b)
        elif kind == "h2":
            sec += 1
            sub = 0
            out.append(("h2", f"{chap}.{sec} {b[1]}" if chap else b[1]))
        elif kind == "h3":
            sub += 1
            out.append(("h3", f"{chap}.{sec}.{sub} {b[1]}"))
        elif kind == "fig":
            fn += 1
            figs[b[1]] = f"{chap}.{fn}"
            out.append(b)
        elif kind == "tab":
            tn += 1
            tabs[b[1]] = f"{chap}.{tn}"
            out.append(b)
        else:
            out.append(b)
    # citations in order of first appearance
    for b in out:
        for text in texts_of(b):
            for m in re.finditer(r"\{ref:([^}]+)\}", text):
                for key in m.group(1).split(","):
                    key = key.strip()
                    if key not in order:
                        order.append(key)
    return out, figs, tabs, order


def texts_of(b):
    kind = b[0]
    if kind in ("p", "note"):
        return [b[1]]
    if kind in ("bul", "num"):
        return list(b[1])
    if kind == "tab":
        return [b[2]] + [c for row in b[4] for c in row]
    if kind == "fig":
        return [b[3]]
    return []


def resolver(figs, tabs, order):
    def resolve(text):
        def fig(m):
            key = m.group(1)
            if key not in figs:
                raise KeyError(f"unknown figure {key}")
            return f"Figure {figs[key]}"

        def tab(m):
            key = m.group(1)
            if key not in tabs:
                raise KeyError(f"unknown table {key}")
            return f"Table {tabs[key]}"

        def ref(m):
            keys = [k.strip() for k in m.group(1).split(",")]
            for k in keys:
                if k not in C.REFERENCES:
                    raise KeyError(f"unknown reference {k}")
            return ", ".join(f"[{order.index(k) + 1}]" for k in keys)

        text = re.sub(r"\{fig:([^}]+)\}", fig, text)
        text = re.sub(r"\{tab:([^}]+)\}", tab, text)
        return re.sub(r"\{ref:([^}]+)\}", ref, text)

    return resolve


def check_references(blocks, figs, tabs, order):
    body = " ".join(t for b in blocks if b[0] in ("p", "bul", "num", "note") for t in texts_of(b))
    body += " " + " ".join(c for b in blocks if b[0] == "tab" for row in b[4] for c in row)
    missing_f = [k for k in figs if f"{{fig:{k}}}" not in body]
    missing_t = [k for k in tabs if f"{{tab:{k}}}" not in body]
    uncited = [k for k in C.REFERENCES if k not in order]
    problems = []
    if missing_f:
        problems.append(f"figures never referenced in text: {missing_f}")
    if missing_t:
        problems.append(f"tables never referenced in text: {missing_t}")
    if uncited:
        problems.append(f"references never cited: {uncited}")
    return problems


# ------------------------------------------------------------------------------------------------
# Block renderers
def para(doc, text, style_name=None, align=None, size=None, bold=None, space_after=None, keep_next=False):
    p = doc.add_paragraph(style=style_name)
    add_rich(p, text, size=size, bold=bold)
    if align is not None:
        p.paragraph_format.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if keep_next:
        p.paragraph_format.keep_with_next = True
    return p


def png_size(path):
    with open(path, "rb") as fh:
        head = fh.read(24)
    if head[:8] == b"\x89PNG\r\n\x1a\n":
        import struct
        return struct.unpack(">II", head[16:24])
    from docx.image.image import Image as DImage
    im = DImage.from_file(str(path))
    return im.px_width, im.px_height


def figure(doc, b, figs, resolve):
    _, key, file, caption, width = b
    path = FIGS / file
    w_px, h_px = png_size(path)
    w = min(width, TEXT_WIDTH)
    if w * h_px / w_px > MAX_FIG_H:
        w = MAX_FIG_H * w_px / h_px
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(w))
    caption_par(doc, "Figure", figs[key], resolve(caption))


def caption_par(doc, label, number, text):
    cp = doc.add_paragraph(style="Caption")
    chap, n = number.split(".")
    r = cp.add_run(f"{label} {chap}.")
    r.bold = True
    rr = cp.add_run()
    fld_char(rr, "begin")
    rr = cp.add_run()
    # \r pins the number: LibreOffice ignores \s 1 (restart at each Heading 1), and without the
    # pin it numbered chapter 5's first figure 5.12.
    instr(rr, f" SEQ {label} \\r {n} \\* ARABIC \\s 1 ")
    rr = cp.add_run()
    fld_char(rr, "separate")
    rr = cp.add_run(n)
    rr.bold = True
    rr = cp.add_run()
    fld_char(rr, "end")
    r = cp.add_run(": ")
    r.bold = True
    add_rich(cp, text)
    if label == "Table":
        cp.paragraph_format.keep_with_next = True
        cp.paragraph_format.space_after = Pt(4)
        cp.paragraph_format.space_before = Pt(10)
    return cp


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)


def table(doc, b, tabs, resolve):
    _, key, caption, header, rows, widths = b
    caption_par(doc, "Table", tabs[key], resolve(caption))
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)
    total = sum(widths)
    inches = [TEXT_WIDTH * w / total for w in widths]
    for ri, values in enumerate([header] + rows):
        row = t.rows[ri]
        if ri == 0:
            trPr = row._tr.get_or_add_trPr()
            h = OxmlElement("w:tblHeader")
            h.set(qn("w:val"), "true")
            trPr.append(h)
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            cell.width = Inches(inches[ci])
            cell.paragraphs[0].style = doc.styles["Table Text"]
            add_rich(cell.paragraphs[0], resolve(value), bold=True if ri == 0 else None)
            if ri == 0:
                set_cell_shading(cell, "E7E9EE")
    # A small gap after the table.
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(4)
    gap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def code_block(doc, b):
    _, title, text = b
    if title:
        p = doc.add_paragraph()
        add_rich(p, title)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
    lines = text.rstrip("\n").split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph(style="Code Block")
        r = p.add_run(line if line else " ")
        para_shading(p, "F3F4F6")
        if i == 0:
            p.paragraph_format.space_before = Pt(4)
        if i == len(lines) - 1:
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = False


def bullets(doc, items, resolve, numbered=False):
    num_id = new_num(NUM_STATE["decimal_abs"], restart=True) if numbered else NUM_STATE["bullet_num"]
    for i, item in enumerate(items):
        p = doc.add_paragraph(style="List Number Custom" if numbered else "List Bullet Custom")
        set_num(p, num_id)
        add_rich(p, resolve(item))
        if i == len(items) - 1:
            p.paragraph_format.space_after = Pt(8)


# ------------------------------------------------------------------------------------------------
# Front matter
def blank(doc, n=1, size=12):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("")
        r.font.size = Pt(size)


def centered(doc, text, size, bold=True, after=6, italic=False, underline=False):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_rich(p, text, size=Pt(size), bold=bold)
    for r in p.runs:
        if italic:
            r.italic = True
        if underline:
            r.underline = True
    return p


def logo_path():
    target = WORK / "mgm-logo.jpeg"
    if not target.exists():
        WORK.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(FORMAT) as z:
            target.write_bytes(z.read("word/media/image1.jpeg"))
    return target


def cover(doc):
    blank(doc, 1)
    centered(doc, "A", 16, after=4)
    centered(doc, f"B.Tech ({C.PH['branch']})", 16, after=4)
    centered(doc, "PROJECT REPORT ON", 16, after=10)
    centered(doc, f"“{C.TITLE}”", 18, after=24)
    centered(doc, "Submitted by", 16, after=10)
    for name, exam in C.STUDENTS:
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(4.3), WD_TAB_ALIGNMENT.LEFT)
        r = p.add_run(f"{name}\t{exam}")
        r.bold = True
        r.font.size = Pt(14)
    blank(doc, 1)
    centered(doc, "Under Guidance of", 16, after=8)
    centered(doc, C.PH["guide"], 16, after=18)
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    p.add_run().add_picture(str(logo_path()), width=Inches(1.9))
    centered(doc, "Institute of Information and Communication Technology (IICT)", 16, after=4)
    centered(doc, "MGM University, Chh. Sambhajinagar", 16, after=4)
    centered(doc, "Academic Year 2026-2027", 16, after=0)


def certificate(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.paragraph_format.space_after = Pt(0)
    blank(doc, 1)
    centered(doc, "Institute of Information and Communication Technology", 16, after=4)
    centered(doc, "MGM University, Chh. Sambhajinagar", 16, after=36)
    centered(doc, "CERTIFICATE", 22, after=36, underline=True)
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(48)
    s = Pt(14)
    names = " and ".join(n for n, _ in C.STUDENTS)
    add_rich(p, f"This is to certify that **{names}** of B.Tech (*{C.PH['branch']}*), Roll No. **{C.PH['roll']}** "
                "has successfully completed Project work for **Multi-disciplinary minor (CYBERSECURITY)** offered by "
                "**Institute of Information and Communication Technology, MGM University**, Chh. Sambhajinagar, under the "
                f"guidance of **{C.PH['guide']}** and submitted the same during the academic year 2026-2027 towards the "
                f"partial fulfillment of degree of B.Tech (*{C.PH['branch']}*).", size=s)
    for left, right, bold in ((C.PH["guide"], "Dr. S. C. Tamane", True), ("Guide", "Director, IICT", True)):
        q = doc.add_paragraph()
        q.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
        q.paragraph_format.space_after = Pt(4)
        q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = q.add_run(f"{left}\t{right}")
        r.bold = bold
        r.font.size = s
    blank(doc, 2)
    q = doc.add_paragraph()
    r = q.add_run(f"Date: {C.PH['date']}")
    r.bold = True
    r.font.size = s


def front_heading(doc, text):
    return doc.add_paragraph(text, style="Front Heading")


def acknowledgement(doc, resolve):
    front_heading(doc, "Acknowledgement")
    for t in C.ACKNOWLEDGEMENT:
        para(doc, resolve(t))
    blank(doc, 2)
    for line in (C.STUDENTS[0][0], f"Roll No. {C.PH['roll']}", C.PH["signature"]):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.bold = True


def toc_field(doc, heading_text, code, entries, style_for):
    """A TOC field whose cached result is `entries` [(level, text, page)]."""
    page_break(doc)
    doc.add_paragraph(heading_text, style="Front Heading" if heading_text != "Contents" else "Contents Heading")
    if not entries:
        entries = [(1, "(no entries)", "")]
    paragraphs = []
    for level, text, page in entries:
        p = doc.add_paragraph(style=style_for(level))
        paragraphs.append((p, text, page))
    for i, (p, text, page) in enumerate(paragraphs):
        if i == 0:
            r = p.add_run()
            fld_char(r, "begin")
            r = p.add_run()
            instr(r, f" {code} ")
            r = p.add_run()
            fld_char(r, "separate")
        p.add_run(text)
        p.add_run("\t")
        p.add_run(str(page))
        if i == len(paragraphs) - 1:
            r = p.add_run()
            fld_char(r, "end")


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------------------------------------
def new_section(doc, fmt, header_text=None):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sectPr = sec._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), "1")
    sectPr.append(pg)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.text = ""
    if header_text:
        hp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(header_text)
        r.italic = True
        r.font.size = Pt(10)
        para_border_bottom(hp)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_simple_field(fp, "PAGE", "1", size=Pt(11))
    return sec


def build(pages):
    """pages: dict with 'toc', 'lof', 'lot' entry lists (or None for placeholders)."""
    blocks, figs, tabs, order = plan(C.BODY)
    problems = check_references(C.BODY, figs, tabs, order)
    if problems:
        raise SystemExit("\n".join(problems))
    resolve = resolver(figs, tabs, order)
    doc = fresh_document()

    # Section 1: cover and certificate (no page numbers).
    cover(doc)
    certificate(doc)

    # Section 2: front matter, roman numerals.
    new_section(doc, "lowerRoman")
    acknowledgement(doc, resolve)
    toc_entries = pages["toc"] if pages else placeholder_toc(blocks)
    toc_field(doc, "Contents", 'TOC \\o "1-2" \\h \\z \\u', toc_entries, lambda lv: f"TOC {lv}")
    page_break(doc)
    front_heading(doc, "Abstract")
    for t in C.ABSTRACT:
        para(doc, resolve(t))
    kw = doc.add_paragraph()
    add_rich(kw, "**Keywords:** " + C.KEYWORDS)
    page_break(doc)
    front_heading(doc, "Abbreviations")
    abb = ("tab", "abbr", "", ["Abbreviation", "Expansion"], [list(x) for x in C.ABBREVIATIONS], [1.2, 4.8])
    abbreviation_table(doc, abb, resolve)
    lof = pages["lof"] if pages else [(1, f"Figure {figs[b[1]]}: {strip_markup(resolve(b[3]))}", "00") for b in blocks if b[0] == "fig"]
    lot = pages["lot"] if pages else [(1, f"Table {tabs[b[1]]}: {strip_markup(resolve(b[2]))}", "00") for b in blocks if b[0] == "tab"]
    toc_field(doc, "List of Figures", 'TOC \\h \\z \\c "Figure"', lof, lambda lv: "Table of Figures")
    toc_field(doc, "List of Tables", 'TOC \\h \\z \\c "Table"', lot, lambda lv: "Table of Figures")

    # Section 3: the report, arabic numerals.
    new_section(doc, "decimal", C.RUNNING_HEADER)
    first = True
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            _, label, title = b
            text = title if label in ("", None) else (f"Chapter {label}: {title}" if label.isdigit() else f"Appendix {label}: {title}")
            h = doc.add_paragraph(text, style="Heading 1")
            if first:
                h.paragraph_format.page_break_before = False
                first = False
        elif kind == "h2":
            doc.add_paragraph(b[1], style="Heading 2")
        elif kind == "h3":
            doc.add_paragraph(b[1], style="Heading 3")
        elif kind == "p":
            para(doc, resolve(b[1]))
        elif kind == "note":
            p = para(doc, resolve(b[1]), size=Pt(11))
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            para_shading(p, "F3F4F6")
        elif kind == "bul":
            bullets(doc, b[1], resolve)
        elif kind == "num":
            bullets(doc, b[1], resolve, numbered=True)
        elif kind == "fig":
            figure(doc, b, figs, resolve)
        elif kind == "tab":
            table(doc, b, tabs, resolve)
        elif kind == "code":
            code_block(doc, b)
        elif kind == "refs":
            for i, key in enumerate(order, 1):
                p = doc.add_paragraph(style="Reference Entry")
                add_rich(p, f"[{i}]\t" + C.REFERENCES[key])
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
        elif kind == "pagebreak":
            page_break(doc)
        else:
            raise ValueError(kind)
    set_core_properties(doc)
    return doc, blocks, figs, tabs, order


def abbreviation_table(doc, b, resolve):
    _, key, caption, header, rows, widths = b
    t = doc.add_table(rows=1 + len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)
    inches = [TEXT_WIDTH * w / sum(widths) for w in widths]
    for ri, values in enumerate([header] + rows):
        for ci, value in enumerate(values):
            cell = t.rows[ri].cells[ci]
            cell.width = Inches(inches[ci])
            cell.paragraphs[0].style = doc.styles["Table Text"]
            add_rich(cell.paragraphs[0], value, bold=True if ri == 0 else None)
            if ri == 0:
                set_cell_shading(cell, "E7E9EE")


def strip_markup(text):
    return re.sub(r"\*\*(.+?)\*\*|`([^`]+)`|\*([^*\s][^*]*?)\*", lambda m: m.group(1) or m.group(2) or m.group(3), text)


def placeholder_toc(blocks):
    entries = [(1, "Acknowledgement", "00"), (1, "Abstract", "00"), (1, "Abbreviations", "00"),
               (1, "List of Figures", "00"), (1, "List of Tables", "00")]
    for b in blocks:
        if b[0] == "h1":
            _, label, title = b
            text = title if not label else (f"Chapter {label}: {title}" if label.isdigit() else f"Appendix {label}: {title}")
            entries.append((1, text, "00"))
        elif b[0] == "h2":
            entries.append((2, b[1], "00"))
    return entries


def set_core_properties(doc):
    cp = doc.core_properties
    cp.title = C.TITLE
    cp.author = C.STUDENTS[0][0]
    cp.subject = "Final Year Project Report, B.Tech, Multi-disciplinary minor (Cybersecurity)"
    cp.keywords = C.KEYWORDS
    cp.comments = ""


# ------------------------------------------------------------------------------------------------
# Page numbers from a LibreOffice render
def render_pdf(docx_path):
    WORK.mkdir(parents=True, exist_ok=True)
    profile = WORK / "lo-profile"
    cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
           "--outdir", str(WORK), str(docx_path)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=600)
    return WORK / (docx_path.stem + ".pdf")


def page_labels(pdf):
    text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, check=True).stdout
    pages = text.split("\f")
    labels = []
    for pg in pages:
        lines = [l.strip() for l in pg.splitlines() if l.strip()]
        last = lines[-1] if lines else ""
        labels.append(last if re.fullmatch(r"[ivxlc]+|\d+", last) else "")
    return pages, labels


def measure(pdf, blocks, figs, tabs):
    from pypdf import PdfReader

    reader = PdfReader(str(pdf))
    pages, labels = page_labels(pdf)
    outline = []

    def walk(items):
        for it in items:
            if isinstance(it, list):
                walk(it)
            else:
                outline.append((it.title.strip(), reader.get_destination_page_number(it)))

    walk(reader.outline)
    used = set()

    def page_of_heading(text):
        for i, (title, pg) in enumerate(outline):
            if i not in used and title == text:
                used.add(i)
                return labels[pg]
        raise SystemExit(f"heading not found in the PDF outline: {text!r}")

    toc = []
    for level, text, _ in placeholder_toc(blocks):
        toc.append((level, text, page_of_heading(text)))

    def caption_pages(prefix_word, mapping, blocks_kind, caption_index):
        out = []
        for b in blocks:
            if b[0] != blocks_kind:
                continue
            num = mapping[b[1]]
            pat = re.compile(rf"^\s*{prefix_word} {re.escape(num)}:")
            found = ""
            for pg, page_text in enumerate(pages):
                if not labels[pg].isdigit():
                    continue
                if any(pat.match(line) for line in page_text.splitlines()):
                    found = labels[pg]
                    break
            if not found:
                raise SystemExit(f"caption not found: {prefix_word} {num}")
            out.append((1, f"{prefix_word} {num}: {strip_markup(resolver(figs, tabs, plan(C.BODY)[3])(b[caption_index]))}", found))
        return out

    lof = caption_pages("Figure", figs, "fig", 3)
    lot = caption_pages("Table", tabs, "tab", 2)
    return {"toc": toc, "lof": lof, "lot": lot}, len(pages) - (1 if not pages[-1].strip() else 0)


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    draft = WORK / "draft.docx"
    doc, blocks, figs, tabs, order = build(None)
    doc.save(str(draft))
    pdf = render_pdf(draft)
    numbers, _ = measure(pdf, blocks, figs, tabs)
    doc, *_ = build(numbers)
    doc.save(str(OUT))
    final_pdf = render_pdf(OUT)
    check, npages = measure(final_pdf, blocks, figs, tabs)
    if check != numbers:
        diff = [(a, b) for k in numbers for a, b in zip(numbers[k], check[k]) if a != b]
        print("page numbers moved between passes; rebuilding once more:", diff[:5])
        doc, *_ = build(check)
        doc.save(str(OUT))
        final_pdf = render_pdf(OUT)
        again, npages = measure(final_pdf, blocks, figs, tabs)
        if again != check:
            raise SystemExit("page numbers did not settle")
        numbers = check
    (WORK / "numbers.json").write_text(json.dumps(numbers, indent=1, ensure_ascii=False))
    print(f"wrote {OUT} ({npages} PDF pages); figures {len(figs)}, tables {len(tabs)}, references {len(order)}")


if __name__ == "__main__":
    main()
